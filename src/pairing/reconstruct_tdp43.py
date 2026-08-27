"""Reconstruct the TDP-43 neurite-localization MPRA intervention pairs.

The source workbook contains oligo identifiers but not oligo sequences.  This
module reproduces the authors' Gencode-vM17/mm10 oligo construction, then uses
the mutation coordinates embedded in the *design identifiers* (never outcome
values) to resolve the one-tile endpoint convention for each gene.

These are multi-base motif-complement interventions, not SNVs.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import re
import urllib.request
from collections import defaultdict
from dataclasses import dataclass, field
from itertools import permutations
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
GTF = ROOT / "data" / "raw" / "gencode.vM17.annotation.gtf.gz"
COUNTS = (
    ROOT
    / "data"
    / "raw"
    / "PMC12864922_supplementary"
    / "44318_2025_653_MOESM4_ESM.xlsx"
)
FIG4E = (
    ROOT
    / "data"
    / "raw"
    / "PMC12864922_figure4"
    / "Figure4"
    / "4E"
    / "Fig4Esource.txt"
)
GENOME_CACHE = ROOT / "data" / "raw" / "tdp43_mm10_gene_slices"
OUT = ROOT / "data" / "processed" / "tdp43_motif_intervention_pairs.csv.gz"
AUDIT = ROOT / "data" / "processed" / "tdp43_pairing_audit.json"

TARGET_GENE_IDS = {
    "ENSMUSG00000004113",
    "ENSMUSG00000022095",
    "ENSMUSG00000028439",
    "ENSMUSG00000029028",
    "ENSMUSG00000029636",
    "ENSMUSG00000031028",
    "ENSMUSG00000032194",
    "ENSMUSG00000034801",
    "ENSMUSG00000035202",
    "ENSMUSG00000035279",
    "ENSMUSG00000038916",
    "ENSMUSG00000042529",
    "ENSMUSG00000043670",
    "ENSMUSG00000058297",
    "ENSMUSG00000061578",
    "ENSMUSG00000090935",
}
MOTIFS = ("GTGTG", "TGTGT", "GTATG")
COMPLEMENT = str.maketrans("ACGT", "TGCA")
ATTR_RE = re.compile(r'(\w+) "([^"]*)"')
BAD_TAGS = {"cds_start_NF", "mRNA_start_NF", "cds_end_NF", "mRNA_end_NF"}
TSL_EXEMPT_NAMES = {"Ksr2", "Klhl8", "Fam120c", "Fam160b2", "Soga3", "Ptp4a2", "Sos2"}


@dataclass
class Transcript:
    gene_id: str
    transcript_id: str
    transcript_type: str
    support_level: str | None
    tags: set[str]
    start: int
    end: int
    exons: list[tuple[int, int]] = field(default_factory=list)
    cds: list[tuple[int, int]] = field(default_factory=list)


@dataclass
class Gene:
    gene_id: str
    name: str
    chrom: str
    strand: str
    start: int
    end: int
    transcripts: list[Transcript] = field(default_factory=list)


@dataclass(frozen=True)
class Oligo:
    pieces: tuple[tuple[int, int], ...]
    kind: str

    @property
    def start(self) -> int:
        return min(a for a, _ in self.pieces)

    @property
    def end(self) -> int:
        return max(b for _, b in self.pieces)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _attrs(text: str) -> dict[str, str]:
    return dict(ATTR_RE.findall(text))


def parse_target_genes() -> dict[str, Gene]:
    genes: dict[str, Gene] = {}
    transcripts: dict[str, Transcript] = {}
    with gzip.open(GTF, "rt") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            fields = line.rstrip().split("\t")
            attrs = _attrs(fields[8])
            gene_id = attrs.get("gene_id", "").split(".")[0]
            if gene_id not in TARGET_GENE_IDS:
                continue
            feature = fields[2]
            if feature == "gene":
                genes[gene_id] = Gene(
                    gene_id=gene_id,
                    name=attrs["gene_name"],
                    chrom=fields[0],
                    strand=fields[6],
                    start=int(fields[3]),
                    end=int(fields[4]),
                )
            elif feature == "transcript":
                transcript_id = attrs["transcript_id"]
                transcript = Transcript(
                    gene_id=gene_id,
                    transcript_id=transcript_id,
                    transcript_type=attrs.get("transcript_type", ""),
                    support_level=attrs.get("transcript_support_level"),
                    tags=set(re.findall(r'tag "([^"]*)"', fields[8])),
                    start=int(fields[3]),
                    end=int(fields[4]),
                )
                transcripts[transcript_id] = transcript
                genes[gene_id].transcripts.append(transcript)
            elif feature in {"exon", "CDS"}:
                transcript = transcripts[attrs["transcript_id"]]
                target = transcript.exons if feature == "exon" else transcript.cds
                target.append((int(fields[3]), int(fields[4])))
    if set(genes) != TARGET_GENE_IDS:
        raise ValueError(f"Missing target genes: {sorted(TARGET_GENE_IDS - set(genes))}")
    return genes


def load_design_ids() -> tuple[set[str], dict[str, str]]:
    workbook = load_workbook(COUNTS, read_only=True, data_only=True)
    sheet = workbook["allcounts"]
    header = [cell.value for cell in next(sheet.iter_rows())]
    oligo_col = header.index("oligo")
    identifiers = {
        str(row[oligo_col])
        for row in sheet.iter_rows(values_only=True)
        if row[oligo_col] and "|" in str(row[oligo_col])
    }
    natural = {identifier for identifier in identifiers if ":" not in identifier}
    mutant_by_base: dict[str, str] = {}
    for identifier in identifiers - natural:
        base = "|".join(identifier.split("|")[:2])
        if base in mutant_by_base:
            raise ValueError(f"Multiple mutant designs for {base}")
        mutant_by_base[base] = identifier
    if len(natural) != 7_389 or len(mutant_by_base) != 4_566:
        raise ValueError(
            f"Unexpected design counts: {len(natural)} natural, {len(mutant_by_base)} mutant"
        )
    if not set(mutant_by_base).issubset(natural):
        raise ValueError("A mutant design lacks its natural counterpart")
    return natural, mutant_by_base


def _positions(intervals: list[tuple[int, int]]) -> list[int]:
    return sorted({position for start, end in intervals for position in range(start, end + 1)})


def _runs(positions: list[int]) -> list[tuple[int, int]]:
    runs: list[list[int]] = []
    for position in positions:
        if not runs or position > runs[-1][1] + 1:
            runs.append([position, position])
        else:
            runs[-1][1] = position
    return [(start, end) for start, end in runs if end > start]


def _merge_like_bedtools(intervals: list[tuple[int, int]]) -> list[tuple[int, int]]:
    """Mimic bedtools merge on the authors' nominally BED-formatted coordinates."""
    merged: list[list[int]] = []
    for start, end in sorted(intervals):
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return [(start, end) for start, end in merged]


def _passes_transcript_filter(gene: Gene, transcript: Transcript) -> bool:
    if transcript.transcript_type != "protein_coding" or transcript.support_level is None:
        return False
    if transcript.tags & BAD_TAGS:
        return False
    tsl = transcript.support_level
    # Preserve the published script literally, including its ENS-ID/name Cacna1b typo.
    return not (
        (tsl == "NA" or int(tsl) > 1)
        and gene.name not in TSL_EXEMPT_NAMES
        and gene.gene_id != "Cacna1b"
    )


def build_coordinate_oligos(gene: Gene, step: int = 6) -> list[Oligo]:
    gene_utr_intervals: list[tuple[int, int]] = []
    junction_blocks: list[list[int]] = []

    for transcript in gene.transcripts:
        if not _passes_transcript_filter(gene, transcript):
            continue
        exon_positions = _positions(transcript.exons)
        cds_positions = _positions(transcript.cds)
        if not cds_positions:
            continue

        if gene.strand == "+":
            utr_start, utr_end = max(cds_positions) + 1, transcript.end
            index = exon_positions.index(max(cds_positions))
            upstream = exon_positions[index - 99 : index + 1]
            utr_length = utr_end - utr_start + 1
            if utr_length < 260:
                extra_needed = 260 - utr_length
                upstream = exon_positions[index - 259 - extra_needed : index + 1]
        else:
            utr_start, utr_end = transcript.start, min(cds_positions) - 1
            index = exon_positions.index(min(cds_positions))
            upstream = exon_positions[index : index + 99]
            utr_length = utr_end - utr_start + 1
            if utr_length < 260:
                extra_needed = 260 - utr_length
                upstream = exon_positions[index : index + 99 + extra_needed]

        utr_positions = sorted(
            set(upstream + list(range(utr_start, utr_end + 1))).intersection(exon_positions)
        )
        if gene.strand == "+":
            last = utr_positions[-1]
            extension = 2_500 if gene.name == "Kcnj12" else 100
            utr_positions += list(range(last + 1, last + extension + 1))
        else:
            last = utr_positions[0]
            extension = 2_901 if gene.name == "Cacna1b" else 100
            utr_positions = list(range(last - extension, last)) + utr_positions

        if len(utr_positions) > 10_000:
            continue
        utr_exons = _runs(utr_positions)
        gene_utr_intervals.extend(utr_exons)

        # The authors' script resets this variable inside each multi-exon
        # transcript; preserving that behavior is necessary for exact recovery.
        if len(utr_exons) > 1:
            junction_blocks = []
            for left, right in zip(utr_exons, utr_exons[1:]):
                last_left, first_right = left[1], right[0]
                if gene.strand == "+" and first_right < utr_start:
                    continue
                if gene.strand == "-" and first_right > utr_end:
                    continue
                left_index = exon_positions.index(last_left)
                right_index = exon_positions.index(first_right)
                flank = 259 - step
                if left_index - flank < 0:
                    missing = flank + 1 - left_index
                    block = (
                        list(range(exon_positions[0] - missing, exon_positions[0]))
                        + exon_positions[: left_index + 1]
                        + exon_positions[right_index : right_index + 259 + step + 1]
                    )
                elif right_index + flank > len(exon_positions) - 1:
                    missing = flank + 1 - len(exon_positions[right_index:])
                    block_end = exon_positions[-1] + missing
                    block = (
                        exon_positions[left_index - flank : left_index + 1]
                        + exon_positions[right_index:]
                        + list(range(exon_positions[-1] + 1, block_end + 1))
                    )
                else:
                    block = exon_positions[
                        left_index - flank : right_index + flank + 1
                    ]
                junction_blocks.append(block)

    regular: list[Oligo] = []
    for start, end in _merge_like_bedtools(gene_utr_intervals):
        if gene.strand == "+":
            current = start
            while current + 259 <= end:
                regular.append(Oligo(((current, current + 259),), "regular_oneexon"))
                current += step
        else:
            current = end
            while current - 259 >= start:
                regular.append(Oligo(((current - 259, current),), "regular_oneexon"))
                current -= step

    junction: list[Oligo] = []
    for block in junction_blocks:
        for index in range(0, len(block), step):
            positions = block[index : index + 260]
            if len(positions) == 260:
                junction.append(Oligo(tuple(_runs(positions)), "junction"))

    oligos = regular + junction
    if gene.strand == "+":
        return sorted(oligos, key=lambda oligo: oligo.start)
    return sorted(oligos, key=lambda oligo: oligo.end, reverse=True)


def fetch_gene_slice(gene: Gene) -> Path:
    GENOME_CACHE.mkdir(parents=True, exist_ok=True)
    # Covers all published special extensions and junction padding.
    start = max(0, gene.start - 4_000)
    end = gene.end + 4_000
    path = GENOME_CACHE / f"{gene.gene_id}.{gene.chrom}.{start}.{end}.json"
    if path.exists():
        return path
    url = (
        "https://api.genome.ucsc.edu/getData/sequence"
        f"?genome=mm10;chrom={gene.chrom};start={start};end={end}"
    )
    with urllib.request.urlopen(url, timeout=120) as response:
        payload = response.read()
    path.write_bytes(payload)
    return path


def load_gene_slice(gene: Gene, fetch: bool) -> tuple[int, str, Path]:
    expected = list(GENOME_CACHE.glob(f"{gene.gene_id}.{gene.chrom}.*.json"))
    path = fetch_gene_slice(gene) if fetch else (expected[0] if len(expected) == 1 else None)
    if path is None or not path.exists():
        raise FileNotFoundError(f"Missing cached mm10 sequence for {gene.gene_id}; rerun with --fetch")
    payload = json.loads(path.read_text())
    return int(payload["start"]), str(payload["dna"]).upper(), path


def oligo_sequence(oligo: Oligo, strand: str, slice_start: int, dna: str) -> str:
    pieces = []
    for start, end in oligo.pieces:
        # UCSC is zero-based half-open; Gencode/author coordinates are one-based inclusive.
        pieces.append(dna[start - 1 - slice_start : end - slice_start])
    if strand == "+":
        sequence = "".join(pieces)
    else:
        sequence = "".join(piece.translate(COMPLEMENT)[::-1] for piece in reversed(pieces))
    if len(sequence) != 260 or set(sequence) - set("ACGTN"):
        raise ValueError(f"Invalid reconstructed oligo sequence of length {len(sequence)}")
    return sequence


def motif_suffix(sequence: str) -> str | None:
    covered: set[int] = set()
    for motif in MOTIFS:
        for match in re.finditer(f"(?={motif})", sequence):
            covered.update(range(match.start(), match.start() + len(motif)))
    if not covered:
        return None
    chunks = _runs(sorted(covered))
    return ";".join(f"{start + 1}:{end + 1}" for start, end in chunks)


def mutate_by_suffix(sequence: str, suffix: str) -> str:
    chars = list(sequence)
    for chunk in suffix.split(";"):
        start, end = map(int, chunk.split(":"))
        for index in range(start - 1, end):
            chars[index] = chars[index].translate(COMPLEMENT)
    return "".join(chars)


def _candidate_alignments(
    gene: Gene,
    coordinate_oligos: list[Oligo],
    expected_count: int,
    observed_suffixes: dict[str, str],
    slice_start: int,
    dna: str,
) -> tuple[list[Oligo], int, dict[str, object]]:
    difference = len(coordinate_oligos) - expected_count
    if difference < 0 or difference > 2:
        raise ValueError(
            f"{gene.name}: generated {len(coordinate_oligos)}, expected {expected_count}"
        )
    candidates: list[tuple[int, int, int, int, list[Oligo], tuple[str, ...]]] = []
    for trim_left in range(difference + 1):
        trim_right = difference - trim_left
        stop = len(coordinate_oligos) - trim_right if trim_right else None
        candidate = coordinate_oligos[trim_left:stop]
        for coordinate_shift in range(-12, 13):
            matches = 0
            compared = 0
            observed_sequences: list[str] = []
            for number, oligo in enumerate(candidate, start=1):
                base_id = f"{gene.gene_id}.{number}|{gene.name}"
                if base_id not in observed_suffixes:
                    continue
                compared += 1
                shifted = Oligo(
                    tuple(
                        (start + coordinate_shift, end + coordinate_shift)
                        for start, end in oligo.pieces
                    ),
                    oligo.kind,
                )
                sequence = oligo_sequence(shifted, gene.strand, slice_start, dna)
                observed_sequences.append(sequence)
                matches += motif_suffix(sequence) == observed_suffixes[base_id]
            candidates.append(
                (
                    matches,
                    compared,
                    trim_left,
                    coordinate_shift,
                    candidate,
                    tuple(observed_sequences),
                )
            )
    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    best_matches, best_compared = candidates[0][:2]
    tied = [candidate for candidate in candidates if candidate[:2] == candidates[0][:2]]
    # Multiple trim/phase parameterizations can produce the same tiled sequences.
    # They are equivalent only if every experimentally mutated parent is identical.
    if len({candidate[5] for candidate in tied}) != 1:
        raise ValueError(f"{gene.name}: design-supported coordinate alignment is ambiguous")
    _, _, trim_left, coordinate_shift, best, _ = min(
        tied, key=lambda item: (abs(item[3]), item[2], item[3])
    )
    shifted_best = [
        Oligo(
            tuple(
                (start + coordinate_shift, end + coordinate_shift)
                for start, end in oligo.pieces
            ),
            oligo.kind,
        )
        for oligo in best
    ]
    shifted_best, reordered_ties = _reconcile_tied_order(
        gene, shifted_best, observed_suffixes, slice_start, dna
    )
    final_matches = 0
    final_compared = 0
    for number, oligo in enumerate(shifted_best, start=1):
        base_id = f"{gene.gene_id}.{number}|{gene.name}"
        if base_id in observed_suffixes:
            final_compared += 1
            final_matches += (
                motif_suffix(oligo_sequence(oligo, gene.strand, slice_start, dna))
                == observed_suffixes[base_id]
            )
    return shifted_best, coordinate_shift, {
        "generated_coordinate_oligos": len(coordinate_oligos),
        "expected_natural_oligos": expected_count,
        "alignment_suffix_matches": final_matches,
        "alignment_suffix_comparisons": final_compared,
        "selected_left_trim": trim_left,
        "selected_right_trim": difference - trim_left,
        "selected_genomic_coordinate_shift_nt": coordinate_shift,
        "equivalent_best_parameterizations": len(tied),
        "design_resolved_tied_order_groups": reordered_ties,
        "candidate_alignment_scores": [
            {
                "matches": matches,
                "comparisons": compared,
                "left_trim": left_trim,
                "coordinate_shift_nt": shift,
            }
            for matches, compared, left_trim, shift, _, _ in candidates[:10]
        ],
    }


def _reconcile_tied_order(
    gene: Gene,
    oligos: list[Oligo],
    observed_suffixes: dict[str, str],
    slice_start: int,
    dna: str,
) -> tuple[list[Oligo], int]:
    """Resolve gffutils ordering ties from encoded, outcome-blind motif ranges."""
    reconciled = list(oligos)
    primary = (lambda oligo: oligo.start) if gene.strand == "+" else (lambda oligo: oligo.end)
    index = 0
    reordered = 0
    while index < len(reconciled):
        stop = index + 1
        while stop < len(reconciled) and primary(reconciled[stop]) == primary(reconciled[index]):
            stop += 1
        group = reconciled[index:stop]
        if len(group) > 1:
            if len(group) > 7:
                raise ValueError(f"{gene.name}: unexpectedly large coordinate-order tie")
            scored: list[tuple[int, tuple[Oligo, ...]]] = []
            for order in permutations(group):
                score = 0
                for offset, oligo in enumerate(order):
                    number = index + offset + 1
                    base_id = f"{gene.gene_id}.{number}|{gene.name}"
                    if base_id in observed_suffixes:
                        score += (
                            motif_suffix(
                                oligo_sequence(oligo, gene.strand, slice_start, dna)
                            )
                            == observed_suffixes[base_id]
                        )
                scored.append((score, order))
            scored.sort(key=lambda item: item[0], reverse=True)
            if len(scored) == 1 or scored[0][0] > scored[1][0]:
                chosen = list(scored[0][1])
                if chosen != group:
                    reconciled[index:stop] = chosen
                    reordered += 1
        index = stop
    return reconciled, reordered


def reconstruct(fetch: bool = False) -> tuple[pd.DataFrame, dict[str, object]]:
    natural_ids, mutant_by_base = load_design_ids()
    genes = parse_target_genes()
    natural_by_gene: dict[str, set[str]] = defaultdict(set)
    for identifier in natural_ids:
        natural_by_gene[identifier.split("|")[1]].add(identifier)
    observed_suffixes = {
        base: mutant_id.rsplit("|", 1)[1] for base, mutant_id in mutant_by_base.items()
    }

    sequences: dict[str, str] = {}
    gene_audits: dict[str, object] = {}
    cache_sources: list[dict[str, object]] = []
    for gene in genes.values():
        slice_start, dna, cache_path = load_gene_slice(gene, fetch)
        coordinate_oligos = build_coordinate_oligos(gene)
        aligned, _, gene_audit = _candidate_alignments(
            gene,
            coordinate_oligos,
            len(natural_by_gene[gene.name]),
            observed_suffixes,
            slice_start,
            dna,
        )
        for number, oligo in enumerate(aligned, start=1):
            identifier = f"{gene.gene_id}.{number}|{gene.name}"
            sequences[identifier] = oligo_sequence(oligo, gene.strand, slice_start, dna)
        reconstructed_ids = {identifier for identifier in sequences if f"|{gene.name}" in identifier}
        if reconstructed_ids != natural_by_gene[gene.name]:
            raise ValueError(f"{gene.name}: reconstructed natural identifier set differs")
        gene_audits[gene.name] = gene_audit
        cache_sources.append(
            {
                "path": cache_path.relative_to(ROOT).as_posix(),
                "sha256": sha256(cache_path),
                "chrom": gene.chrom,
                "start_0based": slice_start,
                "length_nt": len(dna),
            }
        )

    design_failures: list[dict[str, str]] = []
    mutation_records: dict[str, dict[str, str]] = {}
    for base_id, mutant_id in mutant_by_base.items():
        sequence = sequences[base_id]
        reported_suffix = observed_suffixes[base_id]
        predicted_suffix = motif_suffix(sequence)
        if predicted_suffix != reported_suffix:
            design_failures.append(
                {
                    "base_id": base_id,
                    "reported_suffix": reported_suffix,
                    "predicted_suffix": str(predicted_suffix),
                }
            )
            continue
        mutation_records[base_id] = {
            "mutant_id": mutant_id,
            "mutant_sequence": mutate_by_suffix(sequence, reported_suffix),
        }

    figure = pd.read_csv(FIG4E)
    if len(figure) != 9_132 or figure["oligo"].nunique() != 4_566:
        raise ValueError("Unexpected Figure 4E row/pair counts")
    values = figure.pivot(index="oligo", columns="name", values="value")
    meta = figure.groupby("oligo", as_index=True).agg(
        reported_delta=("l2fcdiff", "first"), clip=("CLIP", "first")
    )
    effects = values.join(meta)
    if not ((effects["mutantl2fc"] - effects["wtl2fc"] - effects["reported_delta"]).abs() < 1e-10).all():
        raise ValueError("Figure 4E delta is inconsistent with mutant minus wild type")

    records: list[dict[str, object]] = []
    for source_row, (base_id, row) in enumerate(effects.iterrows(), start=2):
        if base_id not in mutation_records:
            continue
        gene_id_number, gene_name = base_id.split("|")
        gene_id, oligo_number = gene_id_number.rsplit(".", 1)
        suffix = observed_suffixes[base_id]
        changed = [
            index
            for index, (left, right) in enumerate(
                zip(sequences[base_id], mutation_records[base_id]["mutant_sequence"])
            )
            if left != right
        ]
        records.append(
            {
                "dataset": "tdp43_mpra",
                "intervention_type": "multi_base_tdp43_motif_complement",
                "source_figure4e_row": source_row,
                "parent_id": base_id,
                "gene_id": gene_id,
                "gene_name": gene_name,
                "oligo_number": int(oligo_number),
                "parent_sequence": sequences[base_id],
                "edit_ranges_1based_inclusive": suffix,
                "edited_base_count": len(changed),
                "edit_positions_1based": ";".join(str(index + 1) for index in changed),
                "mutant_id": mutation_records[base_id]["mutant_id"],
                "mutant_sequence": mutation_records[base_id]["mutant_sequence"],
                "parent_localization_log2_neurite_soma": float(row["wtl2fc"]),
                "mutant_localization_log2_neurite_soma": float(row["mutantl2fc"]),
                "delta_localization": float(row["reported_delta"]),
                "tdp43_clip_overlap": str(row["clip"]),
            }
        )

    pairs = pd.DataFrame.from_records(records)
    audit: dict[str, object] = {
        "dataset": "TDP-43 neurite-localization MPRA / GSE288185",
        "intervention_type": "multi-base motif-complement replacements; not SNVs",
        "gencode_release": "vM17",
        "genome_assembly": "mm10",
        "source_files": [
            {"path": GTF.relative_to(ROOT).as_posix(), "sha256": sha256(GTF)},
            {"path": COUNTS.relative_to(ROOT).as_posix(), "sha256": sha256(COUNTS)},
            {"path": FIG4E.relative_to(ROOT).as_posix(), "sha256": sha256(FIG4E)},
            *cache_sources,
        ],
        "natural_design_ids": len(natural_ids),
        "mutant_design_ids": len(mutant_by_base),
        "genes": len(genes),
        "exact_motif_suffix_matches": len(mutation_records),
        "design_mismatches_quarantined": len(design_failures),
        "design_mismatch_examples": design_failures[:25],
        "reconstructed_pairs": len(pairs),
        "complete_parent_sequences": int(pairs["parent_sequence"].notna().sum()),
        "complete_mutant_sequences": int(pairs["mutant_sequence"].notna().sum()),
        "complete_outcomes": int(pairs["delta_localization"].notna().sum()),
        "gene_audits": gene_audits,
    }
    return pairs, audit


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fetch", action="store_true", help="Fetch targeted mm10 slices from UCSC")
    args = parser.parse_args()
    pairs, audit = reconstruct(fetch=args.fetch)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    pairs.to_csv(OUT, index=False, compression="gzip")
    AUDIT.write_text(json.dumps(audit, indent=2) + "\n")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
