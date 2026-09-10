"""Independent stability-data inventory and hash-verified raw reporter recovery inputs."""
from __future__ import annotations
import hashlib
import io
import json
import os
from pathlib import Path
from urllib.request import Request,urlopen
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .io import ROOT,sha256,write_json,output_path,load_development
from .forensics import csv

AUDIT=ROOT/'data/external/mechanism_v2/resource_audit'


def validated_resource(label):
    receipts=sorted(AUDIT.glob(label+'_*.bin.json'))
    if len(receipts)!=1:raise ValueError(f'Ambiguous or missing resource: {label}')
    receipt=json.loads(receipts[0].read_text(encoding='utf-8'))
    path=ROOT/receipt['path']
    if sha256(path)!=receipt['sha256']:raise ValueError('External resource hash mismatch')
    return path,receipt


def audit_stability():
    path,receipt=validated_resource('stability_supp1')
    workbook=load_workbook(io.BytesIO(path.read_bytes()),read_only=True,data_only=False)
    if workbook.sheetnames!=['Sup_T2_metaTable']:raise ValueError('Supplement schema changed')
    sheet=workbook['Sup_T2_metaTable']
    values=list(sheet.values)
    formulas=sum(isinstance(v,str) and v.startswith('=') for row in values for v in row)
    if formulas:raise ValueError('Workbook contains formulas; cached-value provenance needs review')
    frame=pd.DataFrame(values[1:],columns=values[0])
    if frame[['Mutant','UTR_Group']].duplicated().any():
        # Preserve overlapping variants; do not choose a first row by convenience.
        duplicates=int(frame[['Mutant','UTR_Group']].duplicated(keep=False).sum())
    else:duplicates=0
    counts={}
    for cell in ['SH','HEK']:
        columns=[f't05_WT_{cell}',f't05_mt_{cell}']
        y=frame[columns].apply(pd.to_numeric,errors='coerce')
        mask=np.isfinite(y).all(axis=1)&y.gt(0).all(axis=1)
        counts[cell]={'positive_finite_pairs':int(mask.sum()),
            'three_prime_positive_finite_pairs':int((mask&frame.UTR_Group.eq("3'UTR")).sum()),
            'three_prime_genes':int(frame.loc[mask&frame.UTR_Group.eq("3'UTR"),'GeneSymbol'].nunique())}
    rows=load_development()
    names=set(rows.gene_name.dropna().astype(str).str.upper())
    overlap=sorted(set(frame.GeneSymbol.dropna().astype(str).str.upper())&names)
    result={'source':receipt,'sheet':'Sup_T2_metaTable','range':f'A1:U{len(values)}',
        'rows':len(frame),'columns':list(frame.columns),'formula_cells':formulas,
        'duplicate_variant_utr_rows':duplicates,'cell_eligibility':counts,
        'development_gene_symbol_overlap':overlap,
        'full_reference_sequences_present':False,'full_mutant_sequences_present':False,
        'alleles_present':True,'half_life_units':'minutes, as reported by time-course assay',
        'proposed_target':'log2(mutant half-life/reference half-life), no significance filter',
        'status':'Not yet admitted to model training: exact assayed reporter sequences must be recovered/verified',
        'raw_recovery_strategy':'Original paired-end HEK early-time-point amplicons; independent external data only',
        'astrocyte_data_accessed':False,'nzip_outcomes_accessed':False}
    write_json('results/mechanism_v2/manifests/stability_supplement_audit.json',result)
    csv('data/interim/mechanism_v2/stability_supplement_table.csv',frame)
    print(json.dumps(result,indent=2),flush=True)
    return result


def download_verified(url,destination,expected_bytes,expected_md5):
    """Resume only the task-owned partial file; publish only after full MD5/size checks."""
    if not url.startswith('https://ftp.sra.ebi.ac.uk/vol1/fastq/SRR222/'):
        raise PermissionError('Only the audited public stability run host/prefix is admitted')
    path=output_path(destination)
    path.parent.mkdir(parents=True,exist_ok=True)
    lock=path.with_suffix(path.suffix+'.lock')
    fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY)
    os.write(fd,str(os.getpid()).encode());os.close(fd)
    try:
        if not path.exists():
            partial=path.with_suffix(path.suffix+'.partial')
            offset=partial.stat().st_size if partial.exists() else 0
            headers={'User-Agent':'RNAddress-independent-stability-reconstruction'}
            if offset:headers['Range']=f'bytes={offset}-'
            with urlopen(Request(url,headers=headers),timeout=45) as response:
                if offset and (response.status!=206 or not response.headers.get('Content-Range','').startswith(f'bytes {offset}-')):
                    raise RuntimeError('Server did not honor resume range; partial retained for diagnosis')
                with partial.open('ab') as stream:
                    while block:=response.read(8*1024*1024):
                        stream.write(block)
                        if stream.tell()>expected_bytes:raise ValueError('Download exceeds declared file size')
                    stream.flush();os.fsync(stream.fileno())
            if partial.stat().st_size!=expected_bytes:raise ValueError('Incomplete download; partial retained')
            digest=hashlib.md5()
            with partial.open('rb') as stream:
                for block in iter(lambda:stream.read(8*1024*1024),b''):digest.update(block)
            if digest.hexdigest()!=expected_md5:raise ValueError('Published ENA MD5 mismatch; partial retained')
            partial.rename(path)
        digest=hashlib.md5()
        with path.open('rb') as stream:
            for block in iter(lambda:stream.read(8*1024*1024),b''):digest.update(block)
        if path.stat().st_size!=expected_bytes or digest.hexdigest()!=expected_md5:
            raise ValueError('Existing raw file differs from published inventory')
        record={'url':url,'path':destination,'bytes':expected_bytes,'md5':expected_md5,'sha256':sha256(path)}
        write_json('results/mechanism_v2/manifests/'+path.name+'.json',record)
        print(f'Raw stability reads verified: {path.name}',flush=True)
        return record
    finally:
        lock.unlink()  # Only the exact task-owned, exclusively created process lock.


def fetch_recovery_reads():
    from concurrent.futures import ThreadPoolExecutor
    path,_=validated_resource('stability_ena_runs')
    runs=pd.read_csv(path,sep='\t')
    chosen=runs[runs.run_accession.eq('SRR22227655')]
    if len(chosen)!=1 or chosen.iloc[0].sample_title!='HEK293_3U_1_30':
        raise ValueError('Expected external early-time-point run not found')
    row=chosen.iloc[0]
    jobs=[]
    for url,digest,size in zip(row.fastq_ftp.split(';'),row.fastq_md5.split(';'),row.fastq_bytes.split(';')):
        name=url.split('/')[-1]
        jobs.append(('https://'+url,'data/external/mechanism_v2/stability_reads/'+name,int(size),digest))
    with ThreadPoolExecutor(max_workers=2) as pool:
        return list(pool.map(lambda args:download_verified(*args),jobs))
