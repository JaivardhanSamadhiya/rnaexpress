"""Primer-only metadata audit of the already cached public SRLE supplement.

The whitelist excludes every numerical outcome workbook and all raw reads.
This module never constructs a speculative full transcript or fits a model.
"""
from .common import ROOT, ART, jsave, sha256
import ast
import hashlib
import io
import json
import zipfile
import openpyxl


BASE = ROOT / 'data/external/research_20260921'
OLIGOS = [
    'HBB Primers F', 'HBB Primers R',
    'spliced HBB Primers F', 'spliced HBB Primers R',
    'HBB Homo_Arm Primers F', 'HBB Homo_Arm Primers R',
    'spliced HBB Homo_Arm Primers F', 'spliced HBB Homo_Arm Primers R',
    'N1 linear Primers F', 'N1 linear Primers R',
    'N1-HBB linear Primers F', 'N1-HBB linear Primers R',
    'N1-HBB-spliced linear Primers R',
    'Gibson_6mer_Random_F', 'Gibson_6mer_Random_R',
]


def reverse_complement(sequence):
    return sequence.translate(str.maketrans('ACGTN', 'TGCAN'))[::-1]


def run():
    sources = {}
    for name in ('srle_article', 'srle_epmc_supplement'):
        path = BASE / name
        receipt = json.loads((BASE / (name + '.receipt.json')).read_text())
        assert sha256(path) == receipt['sha256']
        sources[name] = {
            'path': path.relative_to(ROOT).as_posix(),
            'sha256': receipt['sha256'], 'public_source_url': receipt['url'],
            'original_retrieved_utc': receipt['retrieved_utc'],
        }
    primers = {}
    with zipfile.ZipFile(BASE / 'srle_epmc_supplement') as outer:
        nested = outer.read('csbj.0107.f1.zip')
    with zipfile.ZipFile(io.BytesIO(nested)) as inner:
        for member, sheet, requested in [
            ('Supplement Table1.xlsx', 'Table S1', OLIGOS),
            ('Supplement Table2.xlsx', 'Table S2',
             ['HBB_qPCR F', 'HBB_qPCR R', 'eGFP_qPCR F', 'eGFP_qPCR R']),
        ]:
            payload = inner.read(member)
            sources[member] = {'sha256': hashlib.sha256(payload).hexdigest(),
                               'bytes': len(payload), 'sheet': sheet,
                               'scope': 'published primer/oligo metadata only'}
            workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            assert workbook.sheetnames == [sheet]
            rows = list(workbook[sheet].iter_rows(values_only=True))
            assert rows[0][0].startswith(sheet + '. List of')
            for row_number, row in enumerate(rows, 1):
                name_column, seq_column = (1, 2) if sheet == 'Table S1' else (0, 1)
                name = row[name_column]
                if name in requested:
                    assert name not in primers
                    sequence = row[seq_column]
                    assert isinstance(sequence, str) and not set(sequence) - set('ACGTN')
                    primers[name] = {
                        'sequence_5prime_to_3prime_dna': sequence,
                        'source_member': member, 'source_sheet': sheet, 'excel_row': row_number,
                        'author_description': row[3] if sheet == 'Table S1' else 'RT-qPCR primer',
                    }
            workbook.close()
            assert set(requested).issubset(primers)
    forward = primers['Gibson_6mer_Random_F']['sequence_5prime_to_3prime_dna']
    reverse = primers['Gibson_6mer_Random_R']['sequence_5prime_to_3prime_dna']
    assert reverse_complement(forward) == reverse
    left, right = forward.split('NNNNNN')
    assert len(left) == len(right) == 20 and len(forward) == 46
    genome_reverse = primers['HBB Primers R']['sequence_5prime_to_3prime_dna']
    downstream_forward = primers['N1-HBB linear Primers F']['sequence_5prime_to_3prime_dna']
    assert reverse_complement(genome_reverse).endswith(left)
    assert downstream_forward.startswith(right)
    source = ROOT / 'src/research_20260921/raw_counts.py'
    patterns = None
    for node in ast.parse(source.read_text()).body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'PATTERNS' for t in node.targets):
            patterns = [ast.literal_eval(call.args[0]).decode() for call in node.value.elts]
    assert patterns == ['ATCACTAAGC([ACGT]{6})ATCATAATCA', 'TGATTATGAT([ACGT]{6})GCTTAGTGAT']
    assert left.endswith('ATCACTAAGC') and right.startswith('ATCATAATCA')
    metadata = {
        'status': 'LOCAL_SYNTHESIS_DESIGN_CONTEXT_VERIFIED_FULL_TRANSCRIPT_UNRESOLVED',
        'source_files': sources, 'primer_metadata': primers,
        'local_design': {
            'left_20nt_dna': left, 'randomized_insert_length': 6, 'right_20nt_dna': right,
            'template_dna': forward, 'length_nt': 46,
            'rna_template': forward.replace('T', 'U'),
            'certification_scope': 'author oligo design and overlap orientation; not complete mature reporter RNA',
            'reverse_oligo_complement_match': True,
            'genomic_HBB_reverse_primer_overlap_match': True,
            'downstream_linearization_primer_overlap_match': True,
            'existing_counter_10nt_flank_match': True,
        },
        'counter_source': {'path': source.relative_to(ROOT).as_posix(), 'sha256': sha256(source),
                           'patterns': patterns, 'raw_reads_reopened': False},
        'metadata_correction': {
            'official_screen_cell': 'HEK293T', 'historical_canonical_annotation': 'MCF7',
            'historical_annotation_modified': False,
            'screen_construct': 'article describes intron-containing HBB for subsequent main experiments',
            'per_clone_splice_state_certified': False,
        },
        'unresolved': ['actual amplified donor allele/clone sequence',
                       'complete final backbone and insertion junctions',
                       'experiment-specific transcription start and 3prime cleavage endpoints',
                       'pre-mRNA versus mature isoforms in both measured fractions',
                       'per-clone cryptic splicing/sequence verification'],
        'opened_supplement_members': ['Supplement Table1.xlsx', 'Supplement Table2.xlsx'],
        'outcome_workbooks_opened': False, 'replicate3_read': False,
        'new_downloads': False, 'biological_fitting': False,
    }
    jsave(ART / 'reporter_context_metadata.json', metadata)
    print(json.dumps({'status': metadata['status'], 'primers': len(primers),
                      'local_context_nt': len(forward), 'outcomes_opened': False}, indent=2))


if __name__ == '__main__':
    run()
