"""Inventory archive/sheet structure without calculating any outcome association."""
from .common import ROOT, write_json, write_new
import io
import json
import zipfile
import openpyxl

BASE = ROOT / 'data/external/research_20260921'

def nested_zip(name, member):
    with zipfile.ZipFile(BASE / name) as outer:
        item = outer.getinfo(member)
        if item.file_size > 25 * 1024 * 1024:
            raise ValueError('Unexpected nested archive size')
        return zipfile.ZipFile(io.BytesIO(outer.read(member)))

def run():
    report = {}
    for name, member in [('srle_epmc_supplement', 'csbj.0107.f1.zip'),
                         ('arora_epmc_supplement', 'gkac763_supplemental_files.zip')]:
        archive = nested_zip(name, member)
        inventory = []
        for item in archive.infolist():
            record = {'member': item.filename, 'bytes': item.file_size}
            if item.filename.endswith('.xlsx'):
                wb = openpyxl.load_workbook(io.BytesIO(archive.read(item)), read_only=True, data_only=True)
                record['sheets'] = []
                for ws in wb:
                    # Header/description strings only; no numerical outcomes.
                    headers = [[str(v)[:250] if isinstance(v, str) else ('<number>' if v is not None else '')
                                for v in row] for row in ws.iter_rows(min_row=1, max_row=2, values_only=True)]
                    record['sheets'].append({'name': ws.title, 'rows': ws.max_row,
                                             'columns': ws.max_column, 'first_two_rows': headers})
                wb.close()
            inventory.append(record)
        report[name] = inventory
    write_json(ROOT / 'results/research_20260921/external_asset_schema.json', report)
    print(json.dumps(report, indent=2))

if __name__ == '__main__':
    run()
