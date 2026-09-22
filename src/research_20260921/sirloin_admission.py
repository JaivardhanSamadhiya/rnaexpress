"""Acquire official, freely downloadable SIRLOIN supplements; inspect schema only."""
from .common import ROOT, sha256, write_json
from .resources import URLS, fetch
from concurrent.futures import ThreadPoolExecutor
import json
import sys
import zipfile
import openpyxl

DATA = ROOT/'data/external/research_20260921'
OUT = ROOT/'results/research_20260921'
BASE = 'https://media.springernature.com/original/springer-static/esm/art%3A10.15252%2Fembj.2020106357/MediaObjects/'
ASSETS = {'sirloin_table_ev2.xlsx': BASE+'44318_2021_BFEMBJ2020106357_MOESM3_ESM.xlsx',
          'sirloin_dataset_ev1.xlsx': BASE+'44318_2021_BFEMBJ2020106357_MOESM5_ESM.xlsx'}


def acquire():
    URLS.update(ASSETS)
    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(fetch,ASSETS))


def inspect():
    report = {}
    for name in ASSETS:
        path = DATA/name
        receipt = json.loads((DATA/(name+'.receipt.json')).read_text())
        if receipt['sha256'] != sha256(path):
            raise ValueError('Source file changed')
        with zipfile.ZipFile(path) as archive:
            members = archive.namelist()
            if '[Content_Types].xml' not in members or archive.testzip() is not None:
                raise ValueError('Invalid XLSX archive')
            active = [n for n in members if 'vbaproject' in n.lower() or 'externallinks' in n.lower()]
        workbook = openpyxl.load_workbook(path,read_only=True,data_only=True)
        sheets = []
        for sheet in workbook:
            # Header rows only; do not read outcome rows for schema admission.
            rows = list(sheet.iter_rows(min_row=1,max_row=1,values_only=True))
            sheets.append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,
                           'first_row':rows})
        workbook.close()
        report[name] = {'source':ASSETS[name],'sha256':sha256(path),
                        'archive_crc_valid':True,'active_content_members':active,'sheets':sheets}
    write_json(OUT/'sirloin_schema.json',report)
    print(json.dumps(report,indent=2,default=str),flush=True)


if __name__ == '__main__':
    {'acquire':acquire,'inspect':inspect}[sys.argv[1]]()
