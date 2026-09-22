"""Outcome-free inventory of a new, publicly archived speckle study."""
from .common import ROOT, sha256, write_json, write_new
import io,json,re,zipfile
from collections import Counter
import openpyxl
from rapidfuzz.distance import Levenshtein

def run():
    path=ROOT/'data/external/research_20260921/speckle2026_epmc_supplement'
    if sha256(path)!='059cd16ae5ea92875fa3ef83fcd2c684e660d6548d3530af2403d7b7311a3cf6':
        raise ValueError('Archive changed')
    outer=zipfile.ZipFile(path)
    nested=zipfile.ZipFile(io.BytesIO(outer.read('gkag174_supplemental_files.zip')))
    constructs=zipfile.ZipFile(io.BytesIO(nested.read('Supplementary Data S1.zip')))
    for archive in [outer,nested,constructs]:
        if archive.testzip() is not None: raise ValueError('CRC failure')
    sequences={}; lengths={}; invalid=[]
    for name in constructs.namelist():
        if not name.endswith('.gb'): continue
        content=constructs.read(name).decode('utf-8')
        declared=int(re.search(r'^LOCUS\s+\S+\s+(\d+)\s+bp',content).group(1))
        sequence=re.sub('[^a-zA-Z]','',content.split('ORIGIN',1)[1].split('//')[0]).upper()
        if len(sequence)!=declared or set(sequence)-set('ACGT'): invalid.append(name)
        sequences[name[:-3]]=sequence;lengths[name[:-3]]=declared
    schemas={}
    for name in nested.namelist():
        if not name.endswith('.xlsx'): continue
        content=nested.read(name)
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if z.testzip() is not None: raise ValueError('Workbook CRC failure')
            active=[n for n in z.namelist() if 'vbaproject' in n.lower() or 'externallinks' in n.lower()]
        w=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        schemas[name]={'active_content':active,'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column,
                 'title':next(s.iter_rows(min_row=1,max_row=1,max_col=1,values_only=True))[0]} for s in w]}
        w.close()
    pairs={}
    for a,b in [('M_COLQ_WT','M_COLQ_MUT'),('M_SMN1','M_SMN2')]:
        pairs[a+'__'+b]={'global_plasmid_edit_distance':Levenshtein.distance(sequences[a],sequences[b]),
                        'lengths':[lengths[a],lengths[b]]}
    out=ROOT/'results/research_20260921'
    write_new(out/'speckle2026_plasmids.fasta',''.join('>'+k+'\n'+v+'\n' for k,v in sorted(sequences.items())).encode())
    summary={'archive_sha256':sha256(path),'crc_valid':True,'plasmids':len(sequences),
             'prefix_counts':dict(Counter(n.split('_')[0] for n in sequences)),
             'length_range':[min(lengths.values()),max(lengths.values())],
             'invalid_constructs':invalid,'schemas':schemas,'named_variant_pairs':pairs,
             'outcome_values_read':False,
             'admission':'Useful mechanistic resource; plasmid DNA is not the expressed RNA. Transcript boundaries and construct-to-outcome matching must be established before any prediction. Small set of designed modules is not adequate evidence for broad gene generalization.',
             'license':'Article XML specifies CC BY 4.0; source attribution required.',
             'source':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12962856/'}
    write_json(out/'speckle2026_inventory.json',summary)
    print(json.dumps(summary,indent=2),flush=True)

if __name__=='__main__':run()
